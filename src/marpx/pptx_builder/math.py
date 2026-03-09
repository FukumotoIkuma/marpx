"""Editable math equation rendering for pptx_builder."""

from __future__ import annotations

import logging

from pptx.util import Emu

from marpx.models import UnsupportedElement
from marpx.utils.common import px_to_emu
from marpx.utils.math import latex_to_omml

logger = logging.getLogger(__name__)


def _add_math_equation(slide, element: UnsupportedElement) -> bool:
    """Add an editable math equation to the slide.

    Returns True if successfully added, False if conversion failed
    (caller should fall back to image).
    """
    latex_source = None
    if element.unsupported_info:
        latex_source = element.unsupported_info.latex_source

    if not latex_source:
        return False

    # Convert LaTeX to OMML
    omml_element = latex_to_omml(latex_source)
    if omml_element is None:
        return False

    # Create textbox at the element's position
    left = Emu(px_to_emu(element.box.x))
    top = Emu(px_to_emu(element.box.y))
    width = Emu(px_to_emu(element.box.width))
    height = Emu(px_to_emu(element.box.height))

    txbox = slide.shapes.add_textbox(left, top, width, height)
    tf = txbox.text_frame
    tf.word_wrap = False

    # Get the first paragraph (auto-created) and inject OMML
    p = tf.paragraphs[0]
    p._element.append(omml_element)

    return True
