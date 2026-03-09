"""Table rendering functions for pptx_builder."""

from __future__ import annotations

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu

from marpx.utils.gradient import parse_linear_gradient
from marpx.models import TableCell, TableElement
from marpx.utils.common import px_to_emu

from ._helpers import (
    _build_gradient_fill_xml,
    _remove_existing_fills,
    _set_fill_color,
    _set_srgb_alpha,
)
from .decoration import _add_decoration_shape
from .text import _add_paragraph_runs

# Fill tags relevant to table cell properties (no noFill/pattFill/blipFill).
_CELL_FILL_TAGS: frozenset[str] = frozenset({qn("a:solidFill"), qn("a:gradFill")})


def _remove_table_style(table) -> None:
    """Strip theme table styling so extracted cell styles are authoritative."""
    tbl_pr = table._tbl.tblPr
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is not None:
        tbl_pr.remove(style_id)
    for attr in ("firstRow", "firstCol", "lastRow", "lastCol", "bandRow", "bandCol"):
        if attr in tbl_pr.attrib:
            del tbl_pr.attrib[attr]


def _set_cell_content(pptx_cell, cell: TableCell) -> None:
    """Set content and styling on a PPTX table cell."""
    tf = pptx_cell.text_frame
    tf.margin_top = Emu(px_to_emu(cell.padding.top_px))
    tf.margin_right = Emu(px_to_emu(cell.padding.right_px))
    tf.margin_bottom = Emu(px_to_emu(cell.padding.bottom_px))
    tf.margin_left = Emu(px_to_emu(cell.padding.left_px))
    for p_idx, para in enumerate(cell.paragraphs):
        if p_idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        _add_paragraph_runs(p, para.runs)

    _set_cell_border(pptx_cell, "L", cell.border_left)
    _set_cell_border(pptx_cell, "R", cell.border_right)
    _set_cell_border(pptx_cell, "T", cell.border_top)
    _set_cell_border(pptx_cell, "B", cell.border_bottom)

    if cell.background_gradient:
        _set_cell_gradient_fill(pptx_cell, cell.background_gradient)
    elif cell.background:
        fill = pptx_cell.fill
        _set_fill_color(fill, cell.background)


def _set_cell_gradient_fill(pptx_cell, css_gradient: str) -> None:
    """Apply a linear gradient fill directly to a table cell tcPr node."""
    parsed = parse_linear_gradient(css_gradient)
    if parsed is None:
        return

    tc_pr = pptx_cell._tc.get_or_add_tcPr()
    _remove_existing_fills(tc_pr, fill_tags=_CELL_FILL_TAGS)

    _build_gradient_fill_xml(tc_pr, parsed)


def _configure_table_border_line(line, width_emu: int) -> None:
    """Match PowerPoint's own table border XML shape defaults."""
    line.set("w", str(width_emu))
    line.set("cap", "flat")
    line.set("cmpd", "sng")
    line.set("algn", "ctr")


def _table_border_width_emu(width_px: float) -> int:
    """Map CSS border widths to PowerPoint table border widths."""
    if width_px <= 0:
        return 0
    return max(px_to_emu(width_px), 12700)


def _set_cell_border(pptx_cell, side: str, border) -> None:
    """Apply a single table cell border side via tcPr XML."""
    tc_pr = pptx_cell._tc.get_or_add_tcPr()
    tag = qn(f"a:ln{side}")
    existing = tc_pr.find(tag)
    if existing is not None:
        tc_pr.remove(existing)

    width_emu = _table_border_width_emu(border.width_px)
    if (
        border.width_px <= 0
        or border.style == "none"
        or border.color is None
        or border.color.a <= 0
    ):
        line = etree.SubElement(tc_pr, tag)
        _configure_table_border_line(line, width_emu or 12700)
        etree.SubElement(line, qn("a:noFill"))
        etree.SubElement(line, qn("a:prstDash")).set("val", "solid")
        etree.SubElement(line, qn("a:round"))
        head_end = etree.SubElement(line, qn("a:headEnd"))
        head_end.set("type", "none")
        head_end.set("w", "med")
        head_end.set("len", "med")
        tail_end = etree.SubElement(line, qn("a:tailEnd"))
        tail_end.set("type", "none")
        tail_end.set("w", "med")
        tail_end.set("len", "med")
        return

    line = etree.SubElement(tc_pr, tag)
    _configure_table_border_line(line, width_emu)
    solid_fill = etree.SubElement(line, qn("a:solidFill"))
    srgb = etree.SubElement(solid_fill, qn("a:srgbClr"))
    srgb.set("val", f"{border.color.r:02X}{border.color.g:02X}{border.color.b:02X}")
    _set_srgb_alpha(solid_fill, border.color.a)
    etree.SubElement(line, qn("a:prstDash")).set("val", "solid")
    etree.SubElement(line, qn("a:round"))
    head_end = etree.SubElement(line, qn("a:headEnd"))
    head_end.set("type", "none")
    head_end.set("w", "med")
    head_end.set("len", "med")
    tail_end = etree.SubElement(line, qn("a:tailEnd"))
    tail_end.set("type", "none")
    tail_end.set("w", "med")
    tail_end.set("len", "med")


def _add_table(slide, element: TableElement) -> None:
    """Add a table to the slide, handling colspan and rowspan merges."""
    if not element.table_rows:
        return

    if element.decoration:
        _add_decoration_shape(
            slide,
            element.box,
            element.decoration,
            perspective_px=element.perspective_px,
            element_height_px=element.box.height,
        )

    num_rows = len(element.table_rows)

    # Step 1: Calculate true grid column count by summing colspans per row
    num_cols = 0
    for row in element.table_rows:
        row_cols = sum(cell.colspan for cell in row.cells)
        num_cols = max(num_cols, row_cols)

    if num_cols == 0:
        return

    # Step 2: Create the table shape
    left = Emu(px_to_emu(element.box.x))
    top = Emu(px_to_emu(element.box.y))
    width = Emu(px_to_emu(element.box.width))
    height = Emu(px_to_emu(element.box.height))

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table
    _remove_table_style(table)

    column_widths_px = [0.0] * num_cols

    # Step 3: Build occupancy map and place cells
    occupied: list[list[bool]] = [[False] * num_cols for _ in range(num_rows)]

    # Track merge operations to apply after placement
    merges: list[tuple[int, int, int, int, TableCell]] = []

    for r_idx, row in enumerate(element.table_rows):
        col_cursor = 0
        for cell in row.cells:
            # Find next unoccupied column
            while col_cursor < num_cols and occupied[r_idx][col_cursor]:
                col_cursor += 1

            if col_cursor >= num_cols:
                break

            c_idx = col_cursor

            # Mark occupied cells for this cell's span
            for dr in range(cell.rowspan):
                for dc in range(cell.colspan):
                    if r_idx + dr < num_rows and c_idx + dc < num_cols:
                        occupied[r_idx + dr][c_idx + dc] = True

            # Record merge if needed, otherwise set content directly
            if cell.width_px:
                per_col_width = cell.width_px / max(cell.colspan, 1)
                for dc in range(cell.colspan):
                    if c_idx + dc < num_cols:
                        column_widths_px[c_idx + dc] = max(
                            column_widths_px[c_idx + dc],
                            per_col_width,
                        )

            if cell.rowspan > 1 or cell.colspan > 1:
                merges.append((r_idx, c_idx, cell.rowspan, cell.colspan, cell))
            else:
                _set_cell_content(table.cell(r_idx, c_idx), cell)

            col_cursor = c_idx + cell.colspan

    # Step 4: Apply merges and set content AFTER merging
    for r, c, rowspan, colspan, cell_data in merges:
        end_r = min(r + rowspan - 1, num_rows - 1)
        end_c = min(c + colspan - 1, num_cols - 1)
        if end_r > r or end_c > c:
            table.cell(r, c).merge(table.cell(end_r, end_c))
        _set_cell_content(table.cell(r, c), cell_data)

    if any(width_px > 0 for width_px in column_widths_px):
        total_width_px = sum(width_px for width_px in column_widths_px if width_px > 0)
        if total_width_px > 0:
            for idx, width_px in enumerate(column_widths_px):
                if width_px > 0:
                    table.columns[idx].width = Emu(px_to_emu(width_px))
